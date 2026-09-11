from __future__ import annotations

import json
import math
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "outputs"


def finite(value: object) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(float(value))


def verify_plan_sheet(workbook: openpyxl.Workbook, sheet_name: str) -> None:
    sheet = workbook[sheet_name]
    assert sheet.max_row == 335
    assert sheet.max_column == 147
    for row in (2, 100, 200, 335):
        values = [sheet.cell(row, col).value for col in range(2, 146)]
        total = sheet.cell(row, 146).value
        cost = sheet.cell(row, 147).value
        assert all(finite(value) and value >= 0 for value in values)
        assert finite(total) and abs(sum(values) - total) < 1.0e-5
        assert finite(cost) and cost >= 0


def main() -> None:
    summary = json.loads((OUTPUTS / "analysis_summary.json").read_text(encoding="utf-8"))
    result1 = openpyxl.load_workbook(OUTPUTS / "result1.xlsx", read_only=False, data_only=True)
    plan = result1["计划购电量"]
    total = sum(plan.cell(row, 2).value for row in range(2, 146))
    assert abs(total - summary["problem1"]["total_grid_kwh"]) < 1.0e-5
    result1.close()

    for name, adjusted in (
        ("result2.xlsx", False),
        ("result3.xlsx", True),
        ("result4-2.xlsx", False),
        ("result4-3.xlsx", True),
    ):
        workbook = openpyxl.load_workbook(OUTPUTS / name, read_only=False, data_only=True)
        verify_plan_sheet(workbook, "计划购电量")
        if adjusted:
            verify_plan_sheet(workbook, "调整购电量")
        assert workbook["充放电量"].max_row == 2005
        assert workbook["充放电量"].max_column == 6
        assert workbook["紧急购电量"].max_column == 3
        workbook.close()
        print(f"PASS {name}")

    assert all(
        check.get("passed", check.get("balance_max_abs_kwh", 1.0) < 1.0e-5)
        for check in summary["validation"].values()
    )
    print("PASS model constraints and workbook cross-checks")


if __name__ == "__main__":
    main()
