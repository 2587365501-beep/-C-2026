from __future__ import annotations

import sys
from pathlib import Path

import openpyxl


def main(paths: list[str]) -> None:
    for raw_path in paths:
        path = Path(raw_path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        print(f"\nFILE {path.name}")
        for worksheet in workbook.worksheets:
            print(
                f"  {worksheet.title}: rows={worksheet.max_row}, "
                f"cols={worksheet.max_column}"
            )
            interesting_rows = list(range(1, min(worksheet.max_row, 8) + 1))
            if worksheet.max_row > 10:
                interesting_rows += [worksheet.max_row - 1, worksheet.max_row]
            interesting_cols = list(range(1, min(worksheet.max_column, 10) + 1))
            if worksheet.max_column > 12:
                interesting_cols += [worksheet.max_column - 1, worksheet.max_column]
            for row_index in interesting_rows:
                values = [
                    worksheet.cell(row_index, col_index).value
                    for col_index in interesting_cols
                ]
                print(f"    row {row_index}: {values!r}")
        workbook.close()


if __name__ == "__main__":
    main(sys.argv[1:])
